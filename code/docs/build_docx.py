import os
import sys
import re
from PIL import Image

import docx
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_style_font(style, font_name):
    """Ensure font is mapped to Times New Roman across all XML script families."""
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>')
        rPr.append(rFonts)
    else:
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rFonts.set(qn('w:eastAsia'), font_name)

def add_figure_caption(doc, fig_num, caption_text):
    """Add a figure caption with native Word SEQ field and Caption style."""
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0

    # Prefix 'Hình '
    r1 = p.add_run('Hình ')
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(13)
    r1.italic = True
    r1.font.color.rgb = RGBColor(0, 0, 0)

    # SEQ Hình Field
    fld = parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Hình \\* ARABIC">'
        f'<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
        f'<w:sz w:val="26"/><w:i/><w:color w:val="000000"/></w:rPr>'
        f'<w:t>{fig_num}</w:t></w:r></w:fldSimple>'
    )
    p._p.append(fld)

    # Suffix ': <title>'
    r2 = p.add_run(f': {caption_text}')
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(13)
    r2.italic = True
    r2.font.color.rgb = RGBColor(0, 0, 0)
    return p

def parse_inline(text):
    """Split text into formatted runs: (content, bold, italic, code)."""
    pattern = re.compile(r'(\*\*\*[^*]+?\*\*\*|\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)')
    parts = pattern.split(text)
    runs = []
    for part in parts:
        if not part:
            continue
        if part.startswith('***') and part.endswith('***'):
            runs.append((part[3:-3], True, True, False))
        elif part.startswith('**') and part.endswith('**'):
            runs.append((part[2:-2], True, False, False))
        elif part.startswith('*') and part.endswith('*'):
            runs.append((part[1:-1], False, True, False))
        elif part.startswith('`') and part.endswith('`'):
            runs.append((part[1:-1], False, False, True))
        else:
            runs.append((part, False, False, False))
    return runs

def add_formatted_runs(paragraph, text, default_size=Pt(13), default_bold=False, default_italic=False):
    """Add inline formatted runs to a paragraph with Times New Roman."""
    runs = parse_inline(text)
    for content, bold, italic, code in runs:
        r = paragraph.add_run(content)
        r.font.name = 'Times New Roman'
        r.font.size = default_size
        r.bold = default_bold or bold
        r.italic = default_italic or italic
        r.font.color.rgb = RGBColor(0, 0, 0)
        # Fix font mapping in Word XML
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>')
            rPr.append(rFonts)

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(base_dir, "BAOCAO_DOAN.md")
    diagram_dir = os.path.join(base_dir, "test", "diagram")
    screenshot_dir = os.path.join(base_dir, "test", "screenshots")
    out_file = os.path.join(base_dir, "BAOCAO_DOAN.docx")

    print(f"Reading markdown: {md_file}")
    with open(md_file, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Page setup: A4, Margins: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

        # Footer with page number
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fld = parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="PAGE">'
            f'<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
            f'<w:sz w:val="24"/><w:color w:val="555555"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>'
        )
        p_foot._p.append(fld)

    # Configure Styles
    # Normal
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_style_font(normal_style, 'Times New Roman')

    # Headings
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.line_spacing = 1.3
    h1.paragraph_format.keep_with_next = True
    set_style_font(h1, 'Times New Roman')

    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.line_spacing = 1.3
    h2.paragraph_format.keep_with_next = True
    set_style_font(h2, 'Times New Roman')

    h3 = doc.styles['Heading 3']
    h3.font.name = 'Times New Roman'
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(2)
    h3.paragraph_format.line_spacing = 1.3
    h3.paragraph_format.keep_with_next = True
    set_style_font(h3, 'Times New Roman')

    h4 = doc.styles['Heading 4']
    h4.font.name = 'Times New Roman'
    h4.font.size = Pt(13)
    h4.font.bold = True
    h4.font.italic = True
    h4.font.color.rgb = RGBColor(0, 0, 0)
    h4.paragraph_format.space_before = Pt(6)
    h4.paragraph_format.space_after = Pt(2)
    h4.paragraph_format.line_spacing = 1.3
    h4.paragraph_format.keep_with_next = True
    set_style_font(h4, 'Times New Roman')

    caption_st = doc.styles['Caption']
    caption_st.font.name = 'Times New Roman'
    caption_st.font.size = Pt(13)
    caption_st.font.italic = True
    caption_st.font.color.rgb = RGBColor(0, 0, 0)
    caption_st.paragraph_format.space_before = Pt(3)
    caption_st.paragraph_format.space_after = Pt(6)
    caption_st.paragraph_format.line_spacing = 1.0
    set_style_font(caption_st, 'Times New Roman')

    # Split lines
    lines = md_text.splitlines()
    total_lines = len(lines)
    print(f"Total lines to parse: {total_lines}")

    # Build cover page
    def build_cover_page():
        # Institutional headers
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.line_spacing = 1.2
        p1.paragraph_format.space_before = Pt(10)
        p1.paragraph_format.space_after = Pt(2)
        r = p1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.bold = True

        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.line_spacing = 1.2
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        r = p2.add_run("ĐẠI HỌC UEH - TRƯỜNG CÔNG NGHỆ VÀ THIẾT KẾ")
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True

        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.line_spacing = 1.2
        p3.paragraph_format.space_before = Pt(2)
        p3.paragraph_format.space_after = Pt(2)
        r = p3.add_run("KHOA CÔNG NGHỆ THÔNG TIN KINH DOANH")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.bold = True

        p4 = doc.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p4.paragraph_format.line_spacing = 1.2
        p4.paragraph_format.space_before = Pt(2)
        p4.paragraph_format.space_after = Pt(40)
        r = p4.add_run("CHUYÊN NGÀNH KỸ THUẬT PHẦN MỀM")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.bold = True

        # Thesis Title
        p_t1 = doc.add_paragraph()
        p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t1.paragraph_format.space_before = Pt(20)
        p_t1.paragraph_format.space_after = Pt(10)
        r = p_t1.add_run("KHÓA LUẬN TỐT NGHIỆP")
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.bold = True

        p_t2 = doc.add_paragraph()
        p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t2.paragraph_format.line_spacing = 1.3
        p_t2.paragraph_format.space_before = Pt(10)
        p_t2.paragraph_format.space_after = Pt(60)
        r = p_t2.add_run("THIẾT KẾ VÀ TRIỂN KHAI HỆ THỐNG ERP ĐA CHI NHÁNH THEO KIẾN TRÚC MODULAR MONOLITH")
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.bold = True

        # Student & Supervisor Info
        def add_info_line(label, value):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(1.0)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r_lbl = p.add_run(label)
            r_lbl.font.name = "Times New Roman"
            r_lbl.font.size = Pt(13)
            r_val = p.add_run(value)
            r_val.font.name = "Times New Roman"
            r_val.font.size = Pt(13)
            r_val.bold = True

        add_info_line("Họ tên sinh viên: ", "Nguyễn Quốc Thắng")
        add_info_line("Mã sinh viên: ", "33241022088")
        add_info_line("Lớp: ", "Kỹ thuật phần mềm 01 - VB2K27.2")
        add_info_line("Khóa: ", "27.2")
        add_info_line("Họ tên giáo viên hướng dẫn: ", "Trương Việt Phương")

        # Bottom info
        p_b1 = doc.add_paragraph()
        p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_b1.paragraph_format.space_before = Pt(80)
        p_b1.paragraph_format.space_after = Pt(4)
        r = p_b1.add_run("Niên khóa: 2026")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

        p_b2 = doc.add_paragraph()
        p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_b2.paragraph_format.space_before = Pt(2)
        p_b2.paragraph_format.space_after = Pt(0)
        r = p_b2.add_run("Tp Hồ Chí Minh, ngày…..tháng…..")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

    build_cover_page()

    # Find where content starts: at line with "# LỜI CẢM ƠN"
    start_idx = 0
    for idx, line in enumerate(lines):
        if line.strip().startswith("# LỜI CẢM ƠN"):
            start_idx = idx
            break

    print(f"Content processing starts from line {start_idx + 1}")

    # Regex for image tags
    bt = chr(96)
    re_diagram = re.compile(r'>\s*\*\*\[DIAGRAM — Hình (\d+):\s*([^\]]+)\]\*\*')
    re_screenshot = re.compile(r'>\s*\*\*\[ẢNH CHỤP — Hình (\d+):\s*(.*?)\s*File:\s*' + bt + r'([^' + bt + r']+)' + bt + r'\]\*\*')

    # Regex for headings
    re_h1 = re.compile(r'^#\s+(.+)$')
    re_h2 = re.compile(r'^##\s+(.+)$')
    re_h3 = re.compile(r'^###\s+(.+)$')
    re_h4 = re.compile(r'^####\s+(.+)$')

    # Regex for list items
    re_numbered = re.compile(r'^(\d+)\.\s+(.*)$')
    re_bullet = re.compile(r'^[-*]\s+(.*)$')

    i = start_idx
    inserted_images = 0

    while i < total_lines:
        line = lines[i]
        stripped = line.strip()

        # 1. Skip empty lines (satisfies user requirement: "giữa các dòng không có 1 dòng khoảng trống vô nghĩa")
        if not stripped:
            i += 1
            continue

        # 2. Check for Mermaid code blocks (Sections 4.8.7) -> skip entirely as confirmed by user
        if stripped.startswith("```mermaid"):
            i += 1
            while i < total_lines and not lines[i].strip().startswith("```"):
                i += 1
            i += 1 # skip closing ```
            continue

        # 3. Check for other code blocks (e.g. bash or directory structure)
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < total_lines and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1 # skip closing ```
            
            # Add code paragraph
            p_code = doc.add_paragraph()
            p_code.paragraph_format.left_indent = Inches(0.4)
            p_code.paragraph_format.line_spacing = 1.15
            p_code.paragraph_format.space_before = Pt(4)
            p_code.paragraph_format.space_after = Pt(4)
            r = p_code.add_run("\n".join(code_lines))
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(0, 0, 0)
            continue

        # 4. Check for Tables
        if stripped.startswith("|") and i + 1 < total_lines and "|---" in lines[i+1]:
            table_lines = []
            while i < total_lines and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            # Parse markdown table
            header_line = table_lines[0]
            header_cells = [c.strip() for c in header_line.strip('|').split('|')]
            data_rows = []
            for t_line in table_lines[2:]: # skip separator line
                row_cells = [c.strip() for c in t_line.strip('|').split('|')]
                # normalize length
                while len(row_cells) < len(header_cells):
                    row_cells.append("")
                data_rows.append(row_cells[:len(header_cells)])

            # Create Docx Table
            num_rows = 1 + len(data_rows)
            num_cols = len(header_cells)
            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Header row formatting
            header_tr = table.rows[0]._tr.get_or_add_trPr()
            header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            for c_idx, cell_text in enumerate(header_cells):
                cell = table.cell(0, c_idx)
                # light gray fill
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F0F0"/>')
                cell._tc.get_or_add_tcPr().append(shd)
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, cell_text, default_size=Pt(13), default_bold=True)

            # Data rows formatting
            for r_idx, row_data in enumerate(data_rows):
                tr_pr = table.rows[r_idx + 1]._tr.get_or_add_trPr()
                tr_pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                for c_idx, cell_text in enumerate(row_data):
                    cell = table.cell(r_idx + 1, c_idx)
                    p = cell.paragraphs[0]
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.left_indent = Pt(0)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    add_formatted_runs(p, cell_text, default_size=Pt(13))

            # No extra empty paragraph after table!
            continue

        # 5. Check Diagram placeholder
        m_diag = re_diagram.match(stripped)
        if m_diag:
            fig_num = m_diag.group(1)
            caption_title = m_diag.group(2).strip()
            img_file = os.path.join(diagram_dir, f"hinh{fig_num}.png")

            if os.path.exists(img_file):
                # Calculate display dimensions
                with Image.open(img_file) as im:
                    w_px, h_px = im.size
                aspect = w_px / h_px

                # Fit to printable area (width 15.5cm, max height 20.0cm)
                w_cm = 15.5
                h_cm = w_cm / aspect
                if h_cm > 20.0:
                    h_cm = 20.0
                    w_cm = h_cm * aspect

                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.line_spacing = 1.0
                r = p_img.add_run()
                r.add_picture(img_file, width=Cm(w_cm))

                # Add Caption
                add_figure_caption(doc, fig_num, caption_title)
                inserted_images += 1
            else:
                print(f"WARNING: Diagram file not found: {img_file}")

            i += 1
            continue

        # 6. Check Screenshot placeholder
        m_sc = re_screenshot.match(stripped)
        if m_sc:
            fig_num = m_sc.group(1)
            caption_title = m_sc.group(2).strip()
            fname = m_sc.group(3).strip()
            img_file = os.path.join(screenshot_dir, fname)

            if os.path.exists(img_file):
                with Image.open(img_file) as im:
                    w_px, h_px = im.size
                aspect = w_px / h_px

                w_cm = 15.5
                h_cm = w_cm / aspect
                if h_cm > 20.0:
                    h_cm = 20.0
                    w_cm = h_cm * aspect

                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.line_spacing = 1.0
                r = p_img.add_run()
                r.add_picture(img_file, width=Cm(w_cm))

                # Add Caption
                add_figure_caption(doc, fig_num, caption_title)
                inserted_images += 1
            else:
                print(f"WARNING: Screenshot file not found: {img_file}")

            i += 1
            continue

        # 7. Check Headings
        m_h1 = re_h1.match(stripped)
        if m_h1:
            title_text = m_h1.group(1).strip()
            p = doc.add_paragraph(style='Heading 1')
            # Clean page break before major section/chapter (zero empty paragraph gap)
            p.paragraph_format.page_break_before = True
            add_formatted_runs(p, title_text, default_size=Pt(16), default_bold=True)
            i += 1
            continue

        m_h2 = re_h2.match(stripped)
        if m_h2:
            title_text = m_h2.group(1).strip()
            p = doc.add_paragraph(style='Heading 2')
            add_formatted_runs(p, title_text, default_size=Pt(14), default_bold=True)
            i += 1
            continue

        m_h3 = re_h3.match(stripped)
        if m_h3:
            title_text = m_h3.group(1).strip()
            p = doc.add_paragraph(style='Heading 3')
            add_formatted_runs(p, title_text, default_size=Pt(13), default_bold=True)
            i += 1
            continue

        m_h4 = re_h4.match(stripped)
        if m_h4:
            title_text = m_h4.group(1).strip()
            p = doc.add_paragraph(style='Heading 4')
            add_formatted_runs(p, title_text, default_size=Pt(13), default_bold=True, default_italic=True)
            i += 1
            continue

        # 8. Check Numbered Lists (e.g. 1. 2. 3...)
        m_num = re_numbered.match(stripped)
        if m_num:
            num_str = m_num.group(1)
            content_text = m_num.group(2).strip()

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = -Inches(0.35)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Explicit number run
            r_num = p.add_run(f"{num_str}. ")
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(13)
            r_num.bold = True
            r_num.font.color.rgb = RGBColor(0, 0, 0)

            add_formatted_runs(p, content_text, default_size=Pt(13))
            i += 1
            continue

        # 9. Check Bullet Lists (e.g. - or *)
        m_bullet = re_bullet.match(stripped)
        if m_bullet:
            content_text = m_bullet.group(1).strip()

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = -Inches(0.18)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            r_bullet = p.add_run("• ")
            r_bullet.font.name = "Times New Roman"
            r_bullet.font.size = Pt(13)
            r_bullet.bold = True
            r_bullet.font.color.rgb = RGBColor(0, 0, 0)

            add_formatted_runs(p, content_text, default_size=Pt(13))
            i += 1
            continue

        # 10. Normal Paragraph (or blockquote like > text)
        clean_text = stripped
        is_quote = False
        if clean_text.startswith(">"):
            is_quote = True
            clean_text = clean_text.lstrip(">").strip()

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if is_quote:
            p.paragraph_format.left_indent = Inches(0.3)

        add_formatted_runs(p, clean_text, default_size=Pt(13), default_italic=is_quote)
        i += 1

    print(f"Finished parsing. Total images inserted: {inserted_images}")
    print(f"Saving to {out_file}...")
    doc.save(out_file)
    print("DOCUMENT SAVED SUCCESSFULLY!")

if __name__ == "__main__":
    main()
